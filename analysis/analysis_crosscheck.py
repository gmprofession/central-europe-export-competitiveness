"""Independent pandas cross-check of Central_Europe_Trade_Analysis.xlsx against the raw WDI CSV."""
import sys, numpy as np, pandas as pd
from openpyxl import load_workbook
raw_csv, xlsx = sys.argv[1], sys.argv[2]
ORDER=['HUN','POL','CZE','SVK']
raw=pd.read_csv(raw_csv,dtype=str,keep_default_na=False)
raw=raw[raw['Country Code'].isin(ORDER)]
yrs=[c for c in raw.columns if 'YR' in c]
L=raw.melt(id_vars=['Country Code','Series Code'],value_vars=yrs).assign(Year=lambda x:x.variable.str[:4].astype(int),value=lambda x:pd.to_numeric(x.value))
W=L.pivot_table(index=['Country Code','Year'],columns='Series Code',values='value')
wb=load_workbook(xlsx,data_only=True)
bad=0
def chk(name,a,b,tol=1e-9):
    global bad
    a=np.asarray(a,float); b=np.asarray(b,float)
    ok=np.allclose(a,b,rtol=tol,atol=tol,equal_nan=True)
    if not ok: bad+=1
    print(("OK   " if ok else "FAIL ")+name+("" if ok else f"  maxdiff={np.nanmax(np.abs(a-b))}"))
# clean sheet
ws=wb['03_CLEAN_DATA']; hdr=[c.value for c in ws[1]]
codes=[ws.parent['01_README'][f'C{27+i}'].value for i in range(14)]
rows=[[c.value for c in r] for r in ws.iter_rows(min_row=2,max_row=65)]
X=pd.DataFrame(rows,columns=hdr)
exp=[];
for cc in ORDER:
    for y in range(2010,2026): exp.append([W.loc[(cc,y),c] for c in codes])
chk("03_CLEAN_DATA == raw (64x14)",X.iloc[:,3:].astype(float).values,np.array(exp),1e-12)
assert list(X['Code'])==[c for c in ORDER for _ in range(16)]
# calculations
wc=wb['04_CALCULATIONS']; C=pd.DataFrame([[c.value for c in r] for r in wc.iter_rows(min_row=2,max_row=65,max_col=15)],columns=[c.value for c in wc[1]][:15])
kd=W['NE.EXP.GNFS.KD'].unstack(0)[ORDER]; g=(kd.pct_change()*100).loc[2010:2025]
chk("export growth",C['Export_growth_real_pct'].astype(float).values,np.concatenate([g[c].values for c in ORDER]),1e-9)
tb=((W['NE.EXP.GNFS.CD']-W['NE.IMP.GNFS.CD'])/W['NY.GDP.MKTP.CD']*100).unstack(0)[ORDER].loc[2010:2025]
chk("trade balance %GDP",C['Trade_balance_pct_GDP'].astype(float).values,np.concatenate([tb[c].values for c in ORDER]))
op=((W['NE.EXP.GNFS.CD']+W['NE.IMP.GNFS.CD'])/W['NY.GDP.MKTP.CD']*100).unstack(0)[ORDER].loc[2010:2025]
chk("openness calc",C['Trade_openness_calc_pct'].astype(float).values,np.concatenate([op[c].values for c in ORDER]))
ix=(kd/kd.loc[2010]*100).loc[2010:2025]
chk("export index",C['Export_index_2010=100'].astype(float).values,np.concatenate([ix[c].values for c in ORDER]))
# flags (independent)
def flags(series_by_c,name=''):
    allv=np.concatenate([series_by_c[c] for c in ORDER]); q1,q3=np.percentile(allv,[25,75]); lo,hi=q1-1.5*(q3-q1),q3+1.5*(q3-q1)
    out=[]
    for c in ORDER:
        v=series_by_c[c]; z=(v-v.mean())/v.std(ddof=1)
        for x,zz in zip(v,z): out.append(" ".join(([ "IQR"] if (x<lo or x>hi) else [])+(["Z"] if abs(zz)>2.5 else [])+(["Rebound"] if (name=="Flag_Export_growth" and len(out)%16==0) else [])))
    return out
cols={'Flag_Export_growth':g,'Flag_GDP_growth':W['NY.GDP.MKTP.KD.ZG'].unstack(0)[ORDER].loc[2010:2025],'Flag_FDI_pct_GDP':W['BX.KLT.DINV.WD.GD.ZS'].unstack(0)[ORDER].loc[2010:2025],
      'Flag_Manufacturing':W['NV.IND.MANF.ZS'].unstack(0)[ORDER].loc[2010:2025],'Flag_Trade_balance':tb,'Flag_Trade_openness':W['NE.TRD.GNFS.ZS'].unstack(0)[ORDER].loc[2010:2025]}
for k,df in cols.items():
    mine=flags({c:df[c].values for c in ORDER},k); xl=[(v.strip() if isinstance(v,str) else "") for v in C[k]]
    ok=[a.strip()==b for a,b in zip(mine,xl)]
    if not all(ok): bad+=1
    print(("OK   " if all(ok) else "FAIL ")+k+f"  flagged={sum(1 for m in mine if m)}")
    if not all(ok): print([ (i,mine[i],xl[i]) for i in range(64) if not ok[i]])
# summary
S=wb['06_SUMMARY']
def find(label,col='A',start=1):
    for r in range(start,S.max_row+1):
        if S[f'{col}{r}'].value==label: return r
# stats block spot checks
def statrow(lab,cc):
    for r in range(1,S.max_row+1):
        if S[f'A{r}'].value==lab and S[f'B{r}'].value==cc: return r
for lab,ser in [("Exports % GDP",W['NE.EXP.GNFS.ZS']),("Trade % GDP (openness)",W['NE.TRD.GNFS.ZS']),("FDI net inflows % GDP",W['BX.KLT.DINV.WD.GD.ZS']),("Manufacturing VA % GDP",W['NV.IND.MANF.ZS']),("GDP per capita, current US$",W['NY.GDP.PCAP.CD'])]:
    for cc in ORDER:
        v=ser.loc[cc].loc[2010:2025]; r=statrow(lab,cc)
        chk(f"stats {lab} {cc}",[S[f'{c}{r}'].value for c in 'CDEFGHI'],[v.mean(),v.median(),v.min(),v.max(),v.std(ddof=1),v.loc[2010],v.loc[2025]],1e-9)
for lab,df in [("Export growth, real %",g),("Trade balance % GDP",tb)]:
    for cc in ORDER:
        v=df[cc]; r=statrow(lab,cc)
        chk(f"stats {lab} {cc}",[S[f'{c}{r}'].value for c in 'CDEFGHI'],[v.mean(),v.median(),v.min(),v.max(),v.std(ddof=1),v.loc[2010],v.loc[2025]],1e-9)
# CAGR
r0=find("Country")  # first header 'Country' might be in other tables; locate via CAGR header text
for r in range(1,S.max_row+1):
    if S[f'B{r}'].value=="Exports in year 1": rc=r+1;break
for i,cc in enumerate(ORDER):
    k=kd[cc]; mine=[(k[2025]/k[2010])**(1/15)-1,(k[2019]/k[2010])**(1/9)-1,(k[2025]/k[2019])**(1/6)-1,k[2025]/k[2010]*100]
    chk(f"CAGR {cc}",[S[f'E{rc+i}'].value,S[f'F{rc+i}'].value,S[f'G{rc+i}'].value,S[f'H{rc+i}'].value],mine)
# period averages
def per(lab,df,cc):
    r=statrow(lab,cc); v=df[cc]
    return [S[f'{c}{r}'].value for c in 'CDEFGHI'],[v.loc[2010:2019].mean(),v.loc[2010:2019].median(),v.loc[2020:2021].mean(),v.loc[2020:2021].median(),v.loc[2022:2025].mean(),v.loc[2022:2025].median(),v.loc[2011:2019].mean()]
# stats block and period block share labels; period block rows are later ones -> find second occurrence
def statrow2(lab,cc):
    hits=[r for r in range(1,S.max_row+1) if S[f'A{r}'].value==lab and S[f'B{r}'].value==cc]
    return hits[-1]
for lab,df in [("Export growth, real %",g),("GDP growth %",cols['Flag_GDP_growth']),("Exports % GDP",W['NE.EXP.GNFS.ZS'].unstack(0)[ORDER].loc[2010:2025]),("Trade % GDP",cols['Flag_Trade_openness']),("FDI net inflows % GDP",cols['Flag_FDI_pct_GDP']),("Trade balance % GDP",tb)]:
    for cc in ORDER:
        r=statrow2(lab,cc); v=df[cc]
        chk(f"period {lab} {cc}",[S[f'{c}{r}'].value for c in 'CDEFGHI'],[v.loc[2010:2019].mean(),v.loc[2010:2019].median(),v.loc[2020:2021].mean(),v.loc[2020:2021].median(),v.loc[2022:2025].mean(),v.loc[2022:2025].median(),v.loc[2011:2019].mean()])
# correlations
V={'EG':g,'FDI':cols['Flag_FDI_pct_GDP'],'GDP':cols['Flag_GDP_growth'],'TRD':cols['Flag_Trade_openness'],'MAN':cols['Flag_Manufacturing'],'EXP':W['NE.EXP.GNFS.ZS'].unstack(0)[ORDER].loc[2010:2025]}
import itertools
for r in range(1,S.max_row+1):
    t=S[f'A{r}'].value
    if isinstance(t,str) and ' vs ' in t and len(t)<12:
        a,b=t.split(' vs ')
        pooled=np.corrcoef(np.concatenate([V[a][c] for c in ORDER]),np.concatenate([V[b][c] for c in ORDER]))[0,1]
        wi=[np.corrcoef(V[a][c],V[b][c])[0,1] for c in ORDER]
        exh=np.corrcoef(np.concatenate([V[a][c] for c in ORDER[1:]]),np.concatenate([V[b][c] for c in ORDER[1:]]))[0,1]
        chk(f"corr {t}",[S[f'{c}{r}'].value for c in 'DEFGHI'],[pooled]+wi+[exh])
# print verdicts
for r in range(1,S.max_row+1):
    if S[f'A{r}'].value in('H1','H2','H3','H4') and S[f'B{r}'].value in('SUPPORTED','NOT SUPPORTED'): print(S[f'A{r}'].value,S[f'B{r}'].value)
print("FAILURES:",bad)
