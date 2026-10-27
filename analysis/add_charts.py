"""Phase 4: add the 8 blueprint charts (live Excel charts) and their interpretation to 05_CHARTS.

Usage: python add_charts.py <phase3_workbook.xlsx> <output.xlsx>
Run build_workbook.py first. Every chart reads 03_CLEAN_DATA / 04_CALCULATIONS directly,
so a re-download of the raw CSV refreshes the charts too. Nothing in sheets 02-04 or 06 is changed.
"""
import sys
from openpyxl import load_workbook
from openpyxl.chart import LineChart, BarChart, ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, Alignment, PatternFill

SRC, OUT = sys.argv[1], sys.argv[2]
wb = load_workbook(SRC)
CL, CA, SM = wb["03_CLEAN_DATA"], wb["04_CALCULATIONS"], wb["06_SUMMARY"]
ws = wb["05_CHARTS"]
for rng in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(rng))
ws.delete_rows(1, ws.max_row)

F = "Arial"
TITLE = Font(name=F, size=14, bold=True)
SEC = Font(name=F, size=12, bold=True, color="1F3864")
BOLD = Font(name=F, size=10, bold=True)
TXT = Font(name=F, size=10)
NOTE = Font(name=F, size=9, italic=True, color="595959")
BLUE = Font(name=F, size=10, color="0000FF")
GREEN = Font(name=F, size=10, color="008000")
YFILL = PatternFill("solid", fgColor="FFF2CC")
TOP = Alignment(wrap_text=True, vertical="top")

# Same colour and marker per country in every chart (palette checked for colour-blind separation;
# markers are the second cue so identity never relies on colour alone).
COUNTRIES = [("HUN", "Hungary", "E34948", "circle"),
             ("POL", "Poland", "2A78D6", "square"),
             ("CZE", "Czechia", "008300", "triangle"),
             ("SVK", "Slovakia", "EDA100", "diamond")]
VCOL = {"GDP_growth": "E", "Exports_pct_GDP": "J", "Trade_pct_GDP": "L", "FDI_pct_GDP": "N"}  # 03_CLEAN_DATA
CCOL = {"Export_growth": "D", "Trade_balance": "E", "Export_index": "H"}                     # 04_CALCULATIONS
assert CL["J1"].value == "Exports_pct_GDP" and CL["L1"].value == "Trade_pct_GDP" and CL["N1"].value == "FDI_pct_GDP" and CL["E1"].value == "GDP_growth_pct"
assert CA["D1"].value == "Export_growth_real_pct" and CA["E1"].value == "Trade_balance_pct_GDP" and CA["H1"].value == "Export_index_2010=100"


def blk(i):
    s = 2 + 16 * i
    return s, s + 15


def years_ref():
    return Reference(CL, min_col=3, min_row=2, max_row=17)   # 2010..2025 (same for every country block)


def style_series(s, hexc, marker, line=True):
    if line:
        s.graphicalProperties.line.solidFill = hexc
        s.graphicalProperties.line.width = 22000
    else:
        s.graphicalProperties.line = LineProperties(noFill=True)
    s.marker.symbol = marker
    s.marker.size = 5 if line else 7
    s.marker.graphicalProperties = GraphicalProperties(solidFill=hexc)
    s.marker.graphicalProperties.line.solidFill = hexc
    s.smooth = False


def axes(ch, ytitle, xtitle="Year", ymin=None, ymax=None):
    ch.y_axis.title = ytitle
    ch.x_axis.title = xtitle
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9"))
    ch.x_axis.tickLblPos = "low"          # keep year labels at the bottom when values go negative
    ch.y_axis.number_format = "0"
    ch.y_axis.majorTickMark = "out"
    if ymin is not None: ch.y_axis.scaling.min = ymin
    if ymax is not None: ch.y_axis.scaling.max = ymax
    ch.legend.position = "b"


def line_chart(title, ytitle, sheet, col, countries=COUNTRIES, ymin=None, ymax=None):
    ch = LineChart()
    ch.title = title
    ch.width, ch.height = 17, 9.5
    for code, name, hexc, mk in countries:
        i = [c[0] for c in COUNTRIES].index(code)
        s0, e0 = blk(i)
        cidx = ord(col) - 64
        ser = Series(Reference(sheet, min_col=cidx, min_row=s0, max_row=e0), title=name)
        style_series(ser, hexc, mk)
        ch.series.append(ser)
    ch.set_categories(years_ref())
    axes(ch, ytitle, ymin=ymin, ymax=ymax)
    return ch


row = 1
ws["A1"].value, ws["A1"].font = "05_CHARTS - Phase 4: the 8 blueprint charts and their interpretation", TITLE
ws["A2"].value, ws["A2"].font = ("Every chart is a live Excel chart reading 03_CLEAN_DATA / 04_CALCULATIONS (no pasted values). "
                                 "Same colour and marker per country throughout: Hungary red circle, Poland blue square, Czechia green triangle, "
                                 "Slovakia yellow diamond. Interpretation follows the blueprint: Observation -> Explanation -> Implication. "
                                 "Association language only; no causal claims. Figures quoted in the text are from the WDI release in 02_RAW_DATA."), NOTE
ws.merge_cells("A2:T2"); ws["A2"].alignment = TOP; ws.row_dimensions[2].height = 40

TEXT_COL = "L"           # interpretation sits to the right of each chart (L:T)
BLOCK = 24               # rows per chart block


def write_text(r, num, title, obs, expl, impl, extra=None):
    ws[f"{TEXT_COL}{r}"].value = f"Chart {num}. {title}"; ws[f"{TEXT_COL}{r}"].font = SEC
    parts = [("Observation", obs), ("Explanation", expl), ("Economic implication", impl)]
    if extra: parts.append(extra)
    rr = r + 1
    for lab, body in parts:
        ws[f"{TEXT_COL}{rr}"].value, ws[f"{TEXT_COL}{rr}"].font = lab, BOLD
        rr += 1
        c = ws[f"{TEXT_COL}{rr}"]
        c.value, c.font, c.alignment = body, TXT, TOP
        ws.merge_cells(f"{TEXT_COL}{rr}:T{rr}")
        ws.row_dimensions[rr].height = max(30, 13 * (len(body) // 95 + 1))
        rr += 1
    return rr


def section(r, text):
    ws[f"A{r}"].value, ws[f"A{r}"].font = text, SEC


# ---------------------------------------------------------------- Chart 1
row = 4
section(row, "Chart 1 - Exports of goods and services, % of GDP")
ch = line_chart("Exports, % of GDP", "% of GDP", CL, VCOL["Exports_pct_GDP"], ymin=30, ymax=110)
ws.add_chart(ch, f"A{row + 1}")
write_text(row, 1, "Exports % GDP",
           "Slovakia is the most export-intensive economy: 77% of GDP in 2010, peaking at 99% in 2022 and ending at 85% in 2025. Hungary held 81-87% "
           "for most of the 2010s, then fell to 73% by 2025. Czechia rose to 81% in 2014 (its peak) and drifted back to 67%. Poland is the least "
           "export-intensive but rose steadily from 40% to 50%, peaking at 62% in 2022. Hungary, Poland and Slovakia peak in 2022; Czechia only shows a small 2022 bump.",
           "The shared 2022 spike is consistent with a price effect: energy and general inflation raised the value of traded goods faster than GDP. "
           "Poland's steady rise is consistent with deeper integration into European supply chains. Hungary's decline since 2022 stands out.",
           "This ratio measures trade intensity, not competitiveness. Slovakia is the most exposed to external demand; Poland is catching up from a much larger home market.")

# ---------------------------------------------------------------- Chart 2
row += BLOCK
section(row, "Chart 2 - Trade openness, (exports + imports) % of GDP")
ch = line_chart("Trade openness, % of GDP", "% of GDP", CL, VCOL["Trade_pct_GDP"], ymin=60, ymax=220)
ws.add_chart(ch, f"A{row + 1}")
write_text(row, 2, "Trade openness (WDI Trade % GDP; recomputed and matched in 04_CALCULATIONS)",
           "Mean 2010-2025: Slovakia 178.5%, Hungary 161.8%, Czechia 140.4%, Poland 98.4%. Hungary and Slovakia are above Poland in all 16 years, so H1 is "
           "supported under the frozen rule. Hungary, Poland and Slovakia peak in 2022 (Slovakia 204%); Czechia's peak was 2014 (156%). Since 2022 Hungary "
           "has fallen the most, from 185% to 141%.",
           "Small economies cannot produce everything at home: they import parts, assemble and re-export, so the same value crosses the border twice. "
           "Poland's large domestic market means less activity needs to be traded. The 2022 peak is again consistent with a price effect.",
           "High openness means deep integration into European value chains but also high exposure to external shocks, especially German demand. "
           "Openness above 100% does not mean trade is bigger than the economy: trade is measured in gross value, GDP in value added.")

# ---------------------------------------------------------------- Chart 3 (with 2020 and 2022 shaded)
row += BLOCK
section(row, "Chart 3 - Real GDP growth, % (2020 COVID and 2022 energy shock marked on the axis)")
# helper: year labels that mark the shock years (live formulas; shock years are blue inputs).
# A shaded band needs a bar+line combo chart, which LibreOffice renders incorrectly (it stacks the lines),
# so the Excel chart marks the years on the axis instead; the PNG version (make_charts.py) shades them.
H0 = 320   # helper area, below all charts
ws[f"A{H0 - 2}"].value, ws[f"A{H0 - 2}"].font = "HELPER DATA FOR CHART 3 (axis labels marking the shock years). Blue = input; black = formula.", SEC
ws[f"A{H0 - 1}"].value, ws[f"A{H0 - 1}"].font = "Shock year 1", BOLD
ws[f"B{H0 - 1}"].value, ws[f"B{H0 - 1}"].font, ws[f"B{H0 - 1}"].fill = 2020, BLUE, YFILL
ws[f"C{H0 - 1}"].value, ws[f"C{H0 - 1}"].font, ws[f"C{H0 - 1}"].fill = "COVID", BLUE, YFILL
ws[f"D{H0 - 1}"].value, ws[f"D{H0 - 1}"].font = "Shock year 2", BOLD
ws[f"E{H0 - 1}"].value, ws[f"E{H0 - 1}"].font, ws[f"E{H0 - 1}"].fill = 2022, BLUE, YFILL
ws[f"F{H0 - 1}"].value, ws[f"F{H0 - 1}"].font, ws[f"F{H0 - 1}"].fill = "Energy", BLUE, YFILL
ws[f"A{H0}"].value, ws[f"A{H0}"].font = "Axis label", BOLD
gcol = VCOL["GDP_growth"]
for k in range(16):
    r = H0 + 1 + k
    ws[f"A{r}"].value = (f"=IF('03_CLEAN_DATA'!C{2 + k}=$B${H0 - 1},'03_CLEAN_DATA'!C{2 + k}&\" \"&$C${H0 - 1},"
                         f"IF('03_CLEAN_DATA'!C{2 + k}=$E${H0 - 1},'03_CLEAN_DATA'!C{2 + k}&\" \"&$F${H0 - 1},'03_CLEAN_DATA'!C{2 + k}&\"\"))")
    ws[f"A{r}"].font = TXT
ch = line_chart("Real GDP growth, % (shock years marked)", "% change, constant prices", CL, gcol, ymin=-6, ymax=8)
ch.set_categories(Reference(ws, min_col=1, min_row=H0 + 1, max_row=H0 + 16))
ws.add_chart(ch, f"A{row + 1}")
write_text(row, 3, "GDP growth",
           "In 2020 all four contract: Czechia -5.3%, Hungary -4.3%, Slovakia -2.6%, Poland -2.0%. In 2021 all rebound (Hungary 7.2%, Poland 6.9%). "
           "After 2022 growth slows sharply: in 2023 Hungary is slightly negative (-0.8%) and Czechia and Poland are near zero. Hungary and Czechia also dipped "
           "in 2012-13 (euro-area debt crisis). Poland's only negative year is 2020, and it leads in 2025 at 3.6%.",
           "The most open economies move with European, especially German, demand, so they swing harder. After 2022, high energy prices and high "
           "interest rates weighed most on energy-intensive, trade-dependent economies.",
           "Links to Chart 2: openness brings growth in good years and volatility in bad ones. Poland's large home market appears to act as a buffer.")

# ---------------------------------------------------------------- Chart 4 (two panels)
row += BLOCK
section(row, "Chart 4 - FDI net inflows, % of GDP (Hungary on its own panel)")
ch = line_chart("FDI inflows, % GDP: POL, CZE, SVK", "% of GDP", CL, VCOL["FDI_pct_GDP"], countries=COUNTRIES[1:], ymin=-2, ymax=8)
ch.width = 13
ws.add_chart(ch, f"A{row + 1}")
ch2 = line_chart("FDI inflows, % GDP: Hungary", "% of GDP", CL, VCOL["FDI_pct_GDP"], countries=COUNTRIES[:1], ymin=-60, ymax=120)
ch2.width = 13
ws.add_chart(ch2, f"A{row + 21}")
write_text(row, 4, "FDI % GDP",
           "Poland, Czechia and Slovakia mostly sit between 0% and 6% of GDP. Czechia is the steadiest (mostly 3-5%); Poland peaks at 6.0% in 2022; "
           "Slovakia moves more and dips slightly negative in 2014 and 2020. Hungary is on a different scale: +54% (2016), -40% (2018), +106% (2020), "
           "and -28% to -33% in 2023-2025. 12 of Hungary's 16 years are flagged as outliers in 04_CALCULATIONS.",
           "Swings of this size cannot be factories being built and torn down. They are consistent with pass-through capital: holding companies, "
           "special-purpose entities and intra-group transfers recorded in the balance of payments.",
           "For POL, CZE and SVK, FDI is a stable, modest inflow consistent with value-chain integration. For Hungary the indicator cannot be taken at face "
           "value, which is why the pooled H3 test is dominated by Hungary (see Chart 7).",
           ("Design note", "Hungary is shown on a separate panel with its own scale; on a shared axis it flattens the other three countries."))
row += 20   # extra space for the second panel

# ---------------------------------------------------------------- Chart 5
row += BLOCK
section(row, "Chart 5 - Trade balance, (exports - imports) % of GDP")
ch = line_chart("Trade balance, % of GDP", "% of GDP", CA, CCOL["Trade_balance"], ymin=-8, ymax=10)
ch.y_axis.crosses = "autoZero"
ws.add_chart(ch, f"A{row + 1}")
write_text(row, 5, "Trade balance % GDP (calculated in 04_CALCULATIONS)",
           "Czechia is the most consistent surplus country, mostly 4-8% of GDP, but the surplus almost vanished in 2022 (0.8%). Hungary ran surpluses of "
           "about 2-8% in 2010-2020, slipped into a small deficit in 2021 (-0.4%) and a deficit of -5.3% in 2022. Poland moved from a deficit of -2.4% in 2010 "
           "to steady surpluses (5.6% in 2020, 5.7% in 2023). Slovakia stays close to zero, peaks near 5% in 2012-13, and records the sample's worst deficit "
           "in 2022 (-5.9%). Hungary and Czechia recovered in 2023; Slovakia is still around zero.",
           "2022 is the clear break. All four import nearly all their oil and gas; when energy prices jumped, the import bill rose faster than exports. "
           "Hungary and Slovakia, the most dependent on Russian gas, were hit hardest. The current-account series (06_SUMMARY) shows the same 2022 dip.",
           "A surplus signals resilience, but in very open economies the balance can flip quickly with import prices. Poland's move from deficit to "
           "surplus is one of the more convincing signs of improved competitiveness in the data.")

# ---------------------------------------------------------------- Chart 6
row += BLOCK
section(row, "Chart 6 - Annual real export growth, % (constant 2015 US$)")
ch = line_chart("Real export growth, % per year", "% change", CA, CCOL["Export_growth"], ymin=-10, ymax=20)
ws.add_chart(ch, f"A{row + 1}")
write_text(row, 6, "Annual real export growth",
           "2010 is high everywhere (11% Hungary to 18% Slovakia): the rebound from 2009, tagged 'Rebound'. In 2020 all four fall: Czechia -8.5%, Slovakia -6.4%, "
           "Hungary -6.1%, Poland only -1.1%. In 2021 all bounce back (Poland 12.3%). After that paths split: Poland grows 2-7% a year; Hungary has a strong 2022 "
           "(10.7%) then stalls and is slightly negative in 2024-25; Slovakia is about zero in 2023-24.",
           "This chart strips out prices. In real volumes 2022 was not exceptional for Poland, Czechia or Slovakia, which supports reading the 2022 peaks in "
           "Charts 1-2 as mostly price inflation. Hungary's recent weakness is consistent with weak German demand and a slowdown in car and battery "
           "production, but these data alone cannot show the cause.",
           "Poland is the most consistent real export performer: its only decline is a mild one in 2020. H4 is supported: every country falls in 2020 and "
           "grows in 2021 above its 2010-2019 average.")

# ---------------------------------------------------------------- Chart 7 (scatter, two panels)
row += BLOCK
section(row, "Chart 7 - Real export growth vs FDI % of GDP (scatter, same year)")


def scatter(title, countries, xmin, xmax):
    sc = ScatterChart()
    sc.title, sc.style, sc.scatterStyle = title, 13, "marker"
    sc.width, sc.height = 13, 9.5
    for code, name, hexc, mk in countries:
        i = [c[0] for c in COUNTRIES].index(code)
        s0, e0 = blk(i)
        x = Reference(CL, min_col=ord(VCOL["FDI_pct_GDP"]) - 64, min_row=s0, max_row=e0)
        y = Reference(CA, min_col=ord(CCOL["Export_growth"]) - 64, min_row=s0, max_row=e0)
        ser = Series(y, x, title=name)
        style_series(ser, hexc, mk, line=False)
        sc.series.append(ser)
    axes(sc, "Real export growth, %", "FDI net inflows, % of GDP", ymin=-10, ymax=20)
    sc.x_axis.scaling.min, sc.x_axis.scaling.max = xmin, xmax
    sc.x_axis.majorGridlines = None
    sc.x_axis.number_format = "0"
    return sc


ws.add_chart(scatter("Export growth vs FDI: all four", COUNTRIES, -60, 120), f"A{row + 1}")
ws.add_chart(scatter("Excluding Hungary (supplementary)", COUNTRIES[1:], -2, 8), f"A{row + 21}")

# live correlation box (links to 06_SUMMARY section 6)
eg = next(r for r in range(1, SM.max_row + 1) if SM[f"A{r}"].value == "EG vs FDI")
box = row + 1
ws[f"{TEXT_COL}{box}"].value, ws[f"{TEXT_COL}{box}"].font = "Correlation, FDI % GDP vs real export growth (live from 06_SUMMARY)", BOLD
labels = [("Pooled, n=64 (frozen H3 test)", "D"), ("Hungary, n=16", "E"), ("Poland, n=16", "F"), ("Czechia, n=16", "G"),
          ("Slovakia, n=16", "H"), ("Pooled excl. Hungary, n=48 (supplementary, added after seeing data)", "I")]
for k, (lab, c) in enumerate(labels):
    r = box + 1 + k
    ws[f"{TEXT_COL}{r}"].value, ws[f"{TEXT_COL}{r}"].font = lab, TXT
    cell = ws[f"Q{r}"]
    cell.value, cell.font, cell.number_format = f"='06_SUMMARY'!{c}{eg}", GREEN, "0.00"
write_text(box + 8, 7, "Export growth vs FDI",
           "Pooled correlation is -0.13: essentially no association. Within countries: Poland +0.41, Slovakia +0.35, Czechia +0.10, Hungary -0.28. "
           "Excluding Hungary, the pooled figure turns positive at +0.26. Hungary's points spread across -40% to +106% of GDP while the other three bunch near 0-6%.",
           "Hungary's FDI figures mostly reflect pass-through capital, so they add noise rather than information about real investment. In the other three "
           "the weak positive link is consistent with FDI feeding export capacity. These are same-year correlations; a new plant takes years to export, so a "
           "same-year link would be weak even if FDI matters.",
           "H3 is NOT supported under the frozen rule (pooled correlation is negative), and that is the headline. The ex-Hungary check suggests a modest "
           "association elsewhere, but with 16 points per country it is suggestive only and was added after seeing the data.")
row += 20

# ---------------------------------------------------------------- Chart 8
row += BLOCK + 8
section(row, "Chart 8 - Real exports indexed, 2010 = 100 (constant 2015 US$)")
ch = line_chart("Real exports, 2010 = 100", "Index, 2010 = 100", CA, CCOL["Export_index"], ymin=80, ymax=250)
ws.add_chart(ch, f"A{row + 1}")
write_text(row, 8, "Export index, 2010 = 100",
           "By 2025 Poland reaches 233: real exports more than doubled (CAGR about 5.8% a year). Slovakia ends at 179, Hungary 176, Czechia 173 (CAGR about "
           "3.7-3.9%). Up to 2019 the four tracked fairly closely (Poland 177). The gap opened mostly after 2019: Poland kept climbing while the others flattened, "
           "and Hungary has edged down since 2023 (179 to 176).",
           "Poland's larger, more diversified economy appears less tied to a single sector, while the other three lean heavily on car production and German "
           "demand. The data show the pattern, not the cause: 'consistent with', not 'because of'.",
           "On real export growth, Poland is the clear leader. Together with its move from trade deficit to surplus (Chart 5) and the steadiest GDP growth "
           "(Chart 3), it is the strongest case for the country that strengthened its international competitiveness most. Hungary and Slovakia are the most "
           "open, but openness is not the same as growth.")

ws.column_dimensions["A"].width = 12
for col in "LMNOPQRST":
    ws.column_dimensions[col].width = 13
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- README updates (Phase 4)
R = wb["01_README"]
assert R["A20"].value == "05_CHARTS" and R["A11"].value == "Workbook last updated"
R["B20"].value = "Phase 4: the 8 blueprint charts as live Excel charts (linked to 03/04), each with Observation -> Explanation -> Implication."
R["B11"].value = "25 October 2026 (Phase 4: charts and interpretation added by add_charts.py)"

wb.save(OUT)
print("saved", OUT)
